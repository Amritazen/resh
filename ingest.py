import os
import re
import shutil
from dotenv import load_dotenv
 
load_dotenv()
 
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_chroma import Chroma
from langchain_core.documents import Document
 
def clean_text(text: str) -> str:
    if not text:
        return text
    text = text.encode('utf-8', errors='ignore').decode('utf-8', errors='ignore')
    replacements = {
        '\u2019': "'", '\u2018': "'",
        '\u201c': '"', '\u201d': '"',
        '\u2013': '-', '\u2014': '--',
        '\u2022': '-', '\u00a0': ' ',
        '\u2012': '-', '\u2015': '--',
        '\u00b7': '-', '\u25cf': '-',
        '\u25cb': '-',
    }
    for char, replacement in replacements.items():
        text = text.replace(char, replacement)
    # Remove only unprintable control characters, keep all readable content
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text
 
def load_docx_safe(filepath: str):
    try:
        from langchain_community.document_loaders import Docx2txtLoader
        loader = Docx2txtLoader(filepath)
        return loader.load()
    except Exception as e:
        print(f"Docx2txt failed: {e}, trying python-docx...")
        try:
            import docx
            doc_obj = docx.Document(filepath)
            full_text = [para.text for para in doc_obj.paragraphs if para.text.strip()]
            return [Document(page_content='\n'.join(full_text), metadata={"source": filepath})]
        except Exception as e2:
            print(f"python-docx failed: {e2}")
            return []
 
def get_embeddings():
    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        raise ValueError("GOOGLE_API_KEY not found in environment.")
    return GoogleGenerativeAIEmbeddings(
        model="models/text-embedding-004",
        google_api_key=api_key
    )
 
def ingest_single_file(filepath: str, chroma_path: str):
    print(f"Ingesting: {filepath} -> {chroma_path}")
    ext = filepath.rsplit(".", 1)[-1].lower()
    docs = []
 
    if ext == "pdf":
        from langchain_community.document_loaders import PyPDFLoader
        docs = PyPDFLoader(filepath).load()
    elif ext == "txt":
        try:
            from langchain_community.document_loaders import TextLoader
            docs = TextLoader(filepath, encoding="utf-8").load()
        except Exception:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                docs = [Document(page_content=f.read(), metadata={"source": filepath})]
    elif ext == "docx":
        docs = load_docx_safe(filepath)
    elif ext == "pptx":
        from langchain_community.document_loaders import UnstructuredPowerPointLoader
        docs = UnstructuredPowerPointLoader(filepath).load()
    else:
        print(f"Unsupported: {ext}")
        return
 
    for doc in docs:
        doc.page_content = clean_text(doc.page_content)
 
    docs = [d for d in docs if d.page_content.strip()]
    print(f"Loaded {len(docs)} documents")
 
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=150,
        separators=["\n\n", "\n", ".", " ", ""]
    )
    chunks = splitter.split_documents(docs)
    chunks = [c for c in chunks if c.page_content.strip()]
    print(f"Created {len(chunks)} chunks")
 
    embeddings = get_embeddings()
 
    if os.path.exists(chroma_path):
        try:
            shutil.rmtree(chroma_path)
        except Exception as e:
            print(f"Warning: could not clear old DB: {e}")
 
    Chroma.from_documents(chunks, embeddings, persist_directory=chroma_path)
    print(f"Done. ChromaDB ready at: {chroma_path}")
